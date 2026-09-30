"""The field list the United Kingdom's recording standard asks for.

Read from the template the standard publishes, which is the same kind of source
as the US data dictionary: a statement of what the register asks, not a set of
entries. Measure M5 compares what each register asks and never how much either
holds, so the template is the whole input.

No UK records are downloaded. That was the owner's decision on 30 September
2026, and it keeps this comparison to schema against schema, with no second
body of free text to scrub.

The file is downloaded by script and hashed like every other source. See
governance/data-provenance.md.
"""

from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = REPO_ROOT / "data" / "raw" / "uk_atrs_template_v4.xlsx"

# The sheet that lists every field with its section and the prompt an organisation
# is asked to answer. The workbook's other listing carries broken references.
FIELD_SHEET = "All data"
# The sheet holding the fixed lists of allowed answers.
VALIDATION_SHEET = "Validation"


def fields() -> list[dict]:
    """Every field the standard asks for, in the order the template asks it."""
    import openpyxl

    workbook = openpyxl.load_workbook(TEMPLATE, data_only=True, read_only=True)
    sheet = workbook[FIELD_SHEET]
    out = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        reference, section, field, prompt = row[0], row[1], row[2], row[3]
        if not field:
            continue
        out.append({
            "reference": str(reference).strip() if reference else "",
            "section": str(section).strip() if section else "",
            # A trailing asterisk in the template marks a field asked only under
            # a condition. The marker is not part of the field's name.
            "field": str(field).strip().rstrip("*").strip(),
            "conditional": str(field).strip().endswith("*"),
            "prompt": str(prompt).strip() if prompt else "",
        })
    workbook.close()
    return out


def allowed_answers() -> dict:
    """The fixed lists of answers the standard permits, by column heading.

    These matter to this project more than the field names do. A register that
    publishes its allowed answers has decided in advance what an entry may say,
    which is the thing the US register leaves open.
    """
    import openpyxl

    workbook = openpyxl.load_workbook(TEMPLATE, data_only=True, read_only=True)
    sheet = workbook[VALIDATION_SHEET]
    rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    workbook.close()
    if not rows:
        return {}
    headings = [str(h).strip() if h else "" for h in rows[0]]
    out = {}
    for index, heading in enumerate(headings):
        if not heading:
            continue
        values = []
        for row in rows[1:]:
            if index < len(row) and row[index] is not None:
                value = str(row[index]).strip()
                if value:
                    values.append(value)
        if values:
            out[heading] = values
    return out
